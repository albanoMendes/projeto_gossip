"""Generate docs/slides.pptx from the seminar outline.

This produces an editable PowerPoint deck covering the four required sections
(intro, protocol, guarantees/failure model, demo). Regenerate with:

    .venv/bin/python scripts/make_slides.py

Requires python-pptx (listed in requirements-dev or installed in the venv).
Fill in the real member names via the AUTHORS constant below before delivery.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

AUTHORS = "Albano Mendes, Gustavo Cunha, Arthur Bastos"
GROUP = "Grupo 7"

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "slides.pptx"
IMG_DIR = ROOT / "docs" / "img"

# Map a slide title to a diagram PNG embedded on that slide (if it exists).
SLIDE_IMAGES = {
    "Protocolo em alto nível": "seq_push_pull.png",
    "Arquitetura": "arch_asyncio.png",
    "Demonstração ao vivo": "topo_panel.png",
}

# (title, [bullets]) per slide. Mirrors docs/slides_outline.md.
SLIDES: list[tuple[str, list[str]]] = [
    (
        "Mural Distribuído com Gossip Push-Pull (Anti-Entropia)",
        [
            "Trabalho 2 — Difusão epidêmica (gossip)",
            f"Integrantes: {AUTHORS}",
            f"{GROUP} — Sistemas Distribuídos",
        ],
    ),
    (
        "O problema",
        [
            "Mural de avisos replicado entre 6 processos (A–F), sem servidor central",
            "Um aviso criado em qualquer nó deve chegar aos demais",
            "Falhas reais: mensagens se perdem, processos param",
            "Como espalhar informação de forma simples, robusta e escalável?",
        ],
    ),
    (
        "Introdução: gossip / difusão epidêmica",
        [
            "Troca probabilística e repetida entre pares; espalha como epidemia",
            "Atinge todos os nós em O(log N) rounds (alta probabilidade)",
            "Simples, escalável, robusto, convergente, descentralizado",
            "Parâmetros: fanout (pares por ciclo) e ciclo/intervalo",
        ],
    ),
    (
        "SI (anti-entropia) vs. SIR (rumor-mongering)",
        [
            "SI (usamos): troca periódica e indefinida; mescla conjuntos; reduz entropia; confiável",
            "SI envia o conjunto completo (volume pequeno dispensa otimização)",
            "SIR: propaga 'boatos' por alguns rounds e expira (tombstones); menos banda",
            "SIR pode não alcançar todos; exige controle de rounds — fora do escopo",
        ],
    ),
    (
        "Por que push-pull",
        [
            "push (envio) / pull (peço) / push-pull (ambos)",
            "Ótimo para disseminar rápido e confiável: push na fase inicial, pull na final",
            "Uma troca repara diferenças nos dois sentidos → iniciador e par ficam com a união",
            "Refinamento: proposta original era push-only; push-pull é academicamente mais sólido",
        ],
    ),
    (
        "Protocolo em alto nível",
        [
            "Mensagem JSON sobre UDP: {type, sender, notices}",
            "Ciclo (~300 ms): seleciona 2 pares distintos → 2 trocas GOSSIP→REPLY independentes",
            "GOSSIP: merge + responde REPLY; REPLY: merge",
            "Aviso imutável origem:contador; merge = união por id (dedup natural)",
        ],
    ),
    (
        "Garantias e modelo de falha",
        [
            "Modelo de falha: fail-stop (sem recuperação, estado só em memória) + perda de mensagens",
            "Sem ACK/retransmissão dirigida; reparação por trocas periódicas + redundância (fanout)",
            "Garantia: consistência eventual (merge idempotente/comutativo reduz entropia)",
            "Limites: sem prazo fixo; sem BFT; sem detector de falhas",
        ],
    ),
    (
        "Arquitetura",
        [
            "Módulos: estado / gossip / comunicacao / monitoramento / aplicacao",
            "Um event loop asyncio por nó; NoticeStore único compartilhado",
            "Painel observer read-only em porta de status separada do gossip",
            "Testes: unidade + integração + e2e (CA-01..CA-08) — todos verdes",
        ],
    ),
    (
        "Demonstração ao vivo",
        [
            "1) Subir 6 nós + painel (vazio)",
            "2) post em A → A:1 indo a 6/6",
            "3) post em C → dois avisos convivendo (dedup)",
            "4) loss=0.3 → DESCARTE_SIMULADO e convergência mais lenta, porém completa",
            "5) Ctrl+C em F → painel 'F (indisponível)'; A–E seguem convergindo",
            "6) Encerrar o painel → gossip dos nós continua",
        ],
    ),
    (
        "Conclusão",
        [
            "Gossip push-pull: simples, robusto, convergente; ótimo para disseminação",
            "Trade-offs: consistência eventual, banda redundante, difícil de depurar",
            "Perguntas (cada integrante responde)",
        ],
    ),
]


def build() -> None:
    prs = Presentation()
    title_layout = prs.slide_layouts[0]
    bullet_layout = prs.slide_layouts[1]

    # Title slide.
    title, bullets = SLIDES[0]
    slide = prs.slides.add_slide(title_layout)
    slide.shapes.title.text = title
    slide.placeholders[1].text = "\n".join(bullets)

    # Content slides.
    embedded = 0
    for title, bullets in SLIDES[1:]:
        slide = prs.slides.add_slide(bullet_layout)
        slide.shapes.title.text = title

        image_name = SLIDE_IMAGES.get(title)
        image_path = IMG_DIR / image_name if image_name else None
        has_image = image_path is not None and image_path.exists()

        body_box = slide.placeholders[1]
        if has_image:
            # Narrow the text to the left half, diagram on the right.
            body_box.width = Inches(4.3)
        body = body_box.text_frame
        body.word_wrap = True
        body.clear()
        for i, bullet in enumerate(bullets):
            para = body.paragraphs[0] if i == 0 else body.add_paragraph()
            para.text = bullet
            para.font.size = Pt(16 if has_image else 18)

        if has_image:
            # Place the diagram on the right side, scaled to fit.
            left = Inches(4.8)
            top = Inches(1.8)
            width = Inches(4.9)
            slide.shapes.add_picture(str(image_path), left, top, width=width)
            embedded += 1

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"wrote {OUT} ({len(SLIDES)} slides, {embedded} with embedded diagrams)")


if __name__ == "__main__":
    build()
